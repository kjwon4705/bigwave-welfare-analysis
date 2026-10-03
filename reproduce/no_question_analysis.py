"""모델 B: 추가 질문/대리 추정 없이 전체 표본에 적용하는 일반 안내 규칙.

실제 수급자격 예측 모델이 아니다. 원자료와 기존 모델 A를 수정하지 않는다.
실행: python Code/무질문모델.py (공개 재현본은 --root 로 로컬 프로젝트 지정)
"""
from pathlib import Path
import argparse
import hashlib
import html
import json
import os
import re
import shutil
import sys

import numpy as np
import pandas as pd

VERSION = '2026-10-03-B1'
KEYS = ['가구일련번호', '가구원번호']
STAGES = ['청소년', '청년', '중장년', '노년']
# 수집 원문의 문구를 사람이 검토한 정책별 결정. 기존 제외 정책은 복구하지 않는다.
REVIEW = {
 'WLF00006315': ('포함', '별도 선정기준 없음', '일반가족을 포함하고 별도 선정기준이 없으므로 일반 정보·초기상담 안내만 포함한다. 심화상담/연계서비스 선정은 판정하지 않는다.'),
 'WLF00006295': ('포함', '지원요건 없음', '원문이 모든 주민에 대해 별도 지원요건이 없다고 명시한다. 일반 원격상담 안내만 포함하며 현장방문 가능 여부는 판정하지 않는다.'),
 'WLF00006308': ('제외', '법률상담 필요성 미관측', '국적은 관측되지만 법률 상담이 필요한 상태는 미관측이다. 모든 국민이 상담을 필요로 한다고 가정하지 않는다.'),
 'WLF00004650': ('제외', '대상별 프로그램 범위 미확정', '아동청소년·교사·양육자 등이라는 대상과 개별 교육 프로그램의 범위가 동일하지 않다. 일반 국민 문구만으로 모든 프로그램 자격을 확정하지 않는다.'),
 'WLF00000031': ('제외', '전체 표본에 대한 범위 미확정', '대한민국 국적 응답자에게는 기본 안내가 가능하나 원문에 외국 국적자의 이용 여부가 명시되지 않았다. 외국 국적 556명을 자동 탈락/제거하지 않고 정책을 전체 표본용 모델에서 제외한다.'),
 'WLF00001083': ('제외', '디지털 어려움 미관측', '원문은 AI·디지털에 어려움을 겪는 국민을 대상으로 한다. 연령이나 학력으로 해당 어려움을 추정하지 않는다.'),
 'WLF00000066': ('제외', '교육 필요성·등록 상태 미확정', '정보화교육 필요성이 미관측이다. 복지카드 미소유는 미등록/비장애를 의미하지 않으므로 소유 변수만으로 전체 사람의 여부를 완결하지 않는다.'),
 'WLF00003263': ('제외', '상담·교육활동 필요성 미관측', '초중고 재학은 일부 확인되지만 선정기준의 상담·교육활동 필요성을 관측하지 않았다. 일반학생이라는 표현으로 선정기준을 생략하지 않는다.'),
}
SELECTED = {k for k, v in REVIEW.items() if v[0] == '포함'}


def digest(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def write_json(p, value):
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def csv(df, p):
    df.to_csv(p, index=False, encoding='utf-8-sig')


class NoQuestionGuidanceModel:
    """고정 규칙에 따른 일반 안내. 수급자격은 출력하지 않는다."""
    def __init__(self, spec):
        if spec['version'] != VERSION or spec['output_type'] != 'general_guidance_only':
            raise ValueError('모델 대상 또는 버전이 일치하지 않습니다.')
        if set(spec['policy_ids']) != SELECTED or spec['required_person_features']:
            raise ValueError('미검증 정책 또는 특징량을 사용할 수 없습니다.')
        self.spec = spec

    def predict(self, people, policy_ids=None):
        """입력 행마다 일반 안내1을 반환한다. 미선정 정책을0으로 변환하지 않는다."""
        chosen = self.spec['policy_ids'] if policy_ids is None else list(policy_ids)
        if len(chosen) != len(set(chosen)) or not set(chosen) <= SELECTED:
            raise ValueError('분석 대상 밖 정책입니다. 비해당으로 처리할 수 없습니다.')
        return pd.DataFrame(1, index=people.index, columns=chosen, dtype='int8')


def load_inputs(root):
    config = json.loads((root/'Code/활성정책_설정.json').read_text(encoding='utf-8'))
    policy_path = root/'분석/Raw_data/중앙부처복지서비스_데이터_보완.csv'
    if digest(policy_path) != config['source_sha256']:
        raise ValueError('정책 원문이 변경되었습니다. 재검토 후 모델 버전을 갱신해야 합니다.')
    policies = pd.read_csv(policy_path).fillna('')
    audit_path = root/'분석/Result/정책자격_정밀감사_20261003/전체정책461.csv'
    audit = pd.read_csv(audit_path).fillna('')
    files = list((root/'분석/Raw_data').glob('2025_*.csv'))
    if len(files) != 1:
        raise ValueError('사회조사 원자료를 하나로 특정해야 합니다.')
    people = pd.read_csv(files[0], encoding='cp949', low_memory=False)
    if people.duplicated(KEYS).any() or not people['가구원가중값'].gt(0).all():
        raise ValueError('응답자 키 또는 가중치를 검토해야 합니다.')
    if len(policies) != 461 or set(config['active_policy_ids']) != set(REVIEW):
        raise ValueError('정책 선정 목록이 변경되어 재검토가 필요합니다.')
    excluded = set(config['audit_excluded']) | set(config['retired_model_policies'])
    if SELECTED & excluded or len(excluded) != 453:
        raise ValueError('기존 제외 정책의 재포함 또는 제외 목록 변경')
    if set(audit['서비스ID']) != set(policies['서비스ID']):
        raise ValueError('감사 자료와 정책 원문 불일치')
    return config, policies, audit, people, policy_path, files[0], audit_path


def build(root):
    root = Path(root).resolve()
    c, policies, audit, people, pp, rp, ap = load_inputs(root)
    out = root/'분석/Result/무질문모델_20261003'
    proc = root/'분석/Processed_data/무질문모델'
    site = root/'분석/Result/site'
    for p in [out, proc, out/'charts', site/'tables', site/'charts', site/'reproduce']:
        p.mkdir(parents=True, exist_ok=True)
    spec = {
        'version': VERSION, 'model_type': 'deterministic_constant_guidance',
        'output_type': 'general_guidance_only', 'policy_ids': sorted(SELECTED),
        'required_person_features': [], 'extra_questions': [],
        'population_scope': '사회조사 2025 만13세 이상 응답자 전체; 일반 정보·상담 안내',
        'excluded_outputs': ['실제 선정 여부', '지원금 지급 여부', '서비스 필요성', '현장 이용 가능 여부'],
        'fit_performed': False, 'train_test_split': None,
        'fit_reason': '실제 결과 라벨 없음; 선택된 안내 규칙은 모든 응답자에 대해 상수',
        'policy_source_sha256': digest(pp),
    }
    write_json(proc/'model.json', spec)
    model = NoQuestionGuidanceModel(json.loads((proc/'model.json').read_text(encoding='utf-8')))
    y = model.predict(people)
    csv(pd.concat([people[KEYS], y], axis=1), proc/'개인별_일반안내.csv')
    w = people['가구원가중값']
    review_rows = []
    for _, r in policies[policies['서비스ID'].isin(REVIEW)].iterrows():
        decision, reason, explanation = REVIEW[r['서비스ID']]
        review_rows.append({'서비스ID': r['서비스ID'], '서비스명': r['서비스명'],
            '모델B_선정': decision, '판정근거': reason, '검토설명': explanation,
            '지원대상_보관원문': r['지원대상 상세'], '선정기준_보관원문': r['선정기준'],
            '원문기준연도': r['기준연도'], '출처': r['복지로 상세링크']})
    review = pd.DataFrame(review_rows)
    selected = review[review['모델B_선정'].eq('포함')].copy()
    selected['일반안내_표본수'] = len(people)
    selected['일반안내_가중비율'] = 100.0
    selected['실제선정여부'] = '미관측; 예측하지 않음'
    csv(selected, proc/'선정정책.csv')
    summary = {
        'version': VERSION, 'source_policies': len(policies), 'prior_excluded': 453,
        'reviewed_active': len(review), 'new_excluded': int(review['모델B_선정'].eq('제외').sum()),
        'general_guidance_policies': len(SELECTED), 'all_excluded_from_model_b': 459,
        'validated_actual_eligibility_policies': 0,
        'people': len(people), 'households': int(people['가구일련번호'].nunique()),
        'raw_columns': len(people.columns), 'prediction_features': 0,
        'weighted_mean_guidance': float(np.average(y.sum(axis=1), weights=w)),
        'unweighted_mean_guidance': float(y.sum(axis=1).mean()),
        'guidance_count_variance': float(y.sum(axis=1).var(ddof=0)),
        'fit_performed': False, 'accuracy': None, 'auc': None,
        'foreign_national_respondents_retained': int(people['현재국적코드'].eq(2).sum()),
        'source_sha256': {'policy': digest(pp), 'people': digest(rp), 'previous_audit': digest(ap)},
        'claim_scope': '선정 2개는 일반 안내만. 실제 지원·선정 예측의 검증 완료 정책은 없음. 제외는 부적격이 아니라 모델 범위 밖.',
    }
    age = pd.cut(people['만연령'], [12,18,34,64,np.inf], labels=STAGES)
    if age.isna().any():
        raise ValueError('표본 연령 범위를 재검토해야 합니다.')
    groups = []
    for label in STAGES:
        mask = age.eq(label)
        groups.append({'생애주기': label, '표본수': int(mask.sum()),
            '가중인구비율': float(100*w[mask].sum()/w.sum()),
            '일반안내_가중평균': float(np.average(y.loc[mask].sum(axis=1), weights=w[mask])),
            '일반안내_최소': int(y.loc[mask].sum(axis=1).min()),
            '일반안내_최대': int(y.loc[mask].sum(axis=1).max())})
    groups = pd.DataFrame(groups)
    # 기존 감사는 재분류하지 않고 대표사유별 수만 참고 정보로 공개한다.
    reasons = audit[audit['서비스ID'].isin(c['audit_excluded'])].groupby('대표사유').size().reset_index(name='정책수')
    comparison = pd.DataFrame([
        ['분석 대상 정책수', '8', '2 (일반 안내 한정)'],
        ['선정 원칙', '이용 희망 등을 전제한 기본대상 안내', '전 표본에서 추가 정보가 필요하지 않은 일반 안내만'],
        ['기존 제외 정책 453개', '제외 유지', '제외 유지'],
        ['개인별 미확인', '별도 표시', '미확인이 남는 정책 자체를 모델 범위에서 제외'],
        ['평균의 의미', '기본대상 부합 수 6.02', '일반 안내 수 2.00; 측정 범위가 달라 성능 비교 불가'],
        ['실제 지원 여부 정답', '없음', '없음'],
        ['학습/정확도', '학습하지 않음', '상수 규칙; 학습 및 정확도 보고하지 않음'],
    ], columns=['항목','모델A_기존분석','모델B_무질문분석'])
    flow = pd.DataFrame({'단계':['원정책','기존 제외','이번 검토','이번 추가 제외','일반 안내 포함'],
                         '정책수':[461,453,8,6,2]})
    tables = {'선정검토8개':review, '선정2개':selected, '선정흐름':flow, '생애주기':groups,
              '기존모델비교':comparison, '이전감사제외사유':reasons}
    for name, t in tables.items():
        csv(t, out/(name+'.csv'))
    with pd.ExcelWriter(out/'무질문모델_집계표.xlsx',engine='openpyxl') as writer:
        from openpyxl.styles import Alignment, Font, PatternFill
        for name,t in tables.items():
            t.to_excel(writer,sheet_name=name,index=False)
            ws=writer.book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=48 if any(k in str(col[0].value) for k in ['설명','원문','출처','모델A','모델B']) else 26
            for cell in ws[1]:
                cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='214B62')
            for row in ws.iter_rows():
                for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
            for i in range(2,len(t)+2):ws.row_dimensions[i].height=75 if name=='선정검토8개' else 45
    make_figures(root,out,groups)
    write_json(out/'summary.json',summary)
    write_json(out/'model.json',spec)
    build_page(site,out,summary,review,selected,groups,comparison,reasons)
    for name in ['선정검토8개.csv','선정2개.csv','선정흐름.csv','생애주기.csv','기존모델비교.csv','이전감사제외사유.csv','무질문모델_집계표.xlsx']:
        shutil.copy2(out/name,site/'tables'/('no_question_'+name))
    shutil.copy2(out/'summary.json',site/'reproduce/no_question_summary.json')
    shutil.copy2(out/'model.json',site/'reproduce/no_question_model.json')
    shutil.copy2(Path(__file__),site/'reproduce/no_question_analysis.py')
    for p in (out/'charts').glob('*.png'):shutil.copy2(p,site/'charts'/p.name)
    readme='''# 모델 B: 무질문 일반 안내 규칙

실제 수급자격 예측이 아니다. 선정된 2개 일반 상담·정보 안내만 출력한다.
기존 제외453개를 유지하고 활성8개 원문을 다시 검토해6개를 추가 제외했다.
개인정보를 추정하거나 질문하지 않는다. 실제 행정 결과 라벨은 없고 상수 출력이므로 학습/정확도/AUC를 산출하지 않는다.

## 재현
기존 로컬 프로젝트에서 `python Code/무질문모델.py`를 실행한다.
공개 스크립트는 `python no_question_analysis.py --root "로컬 프로젝트 경로"`로 실행할 수 있다.
필요 입력은 Code/활성정책_설정.json, 분석/Raw_data/중앙부처복지서비스_데이터_보완.csv,
분석/Raw_data/2025_*.csv(사회조사 원자료 한 파일), 분석/Result/정책자격_정밀감사_20261003/전체정책461.csv,
기존 분석으로 생성한 분석/Result/site/index.html이다. pandas, numpy, matplotlib, openpyxl을 사용한다.
원자료는 이용권한이 있는 로컬 환경에만 보관한다. 공개 사이트에는 개인별 결과가 없다.
원정책 해시나 활성목록이 달라지면 중단한다. 개인자료/감사자료 해시는 summary.json에 기록하며 다른 입력에서는 결과를 재검토해야 한다.
NoQuestionGuidanceModel.predict는 일반 안내1만 반환한다. 미선정 정책을 요청하면 예외를 반환하며0(부적격)으로 대체하지 않는다.
현장 상담 여부, 서비스 필요성, 심화 서비스, 최신 접수 여부는 모델이 판정하지 않는다.
'''
    (site/'reproduce/no_question_README.md').write_text(readme,encoding='utf-8')
    base_readme=site/'reproduce/README.md'
    if base_readme.exists():
        txt=base_readme.read_text(encoding='utf-8')
        if 'no_question_analysis.py' not in txt:
            txt+='\n모델 B도 함께 생성하므로 no_question_analysis.py를 Code/무질문모델.py로 저장해야 합니다. 모델 B 추가 입력과 결과 의미는 no_question_README.md를 참고하세요.\n'
            base_readme.write_text(txt,encoding='utf-8')
    home_readme=site/'README.md'
    if home_readme.exists():
        txt=home_readme.read_text(encoding='utf-8')
        if 'no_question_model.html' not in txt:
            txt+='\n## 별도 모델 B\n\n[무질문 모델 보고서](https://kjwon4705.github.io/bigwave-welfare-analysis/no_question_model.html): 기존 데이터만으로 전표본에 안내 가능한 일반 서비스2개. 실제 수급자격 예측이나 새 학습모델이 아니다. 선정 근거·집계표·실행 코드를 함께 공개한다.\n'
            home_readme.write_text(txt,encoding='utf-8')
    report=f'''# 모델 B: 추가 질문 없는 분석

기존 데이터만으로 실제 지원 여부를 판정하고자 했으나 행정 선정/지급 정답이 없으며, 이전 제외453개를 유지한 후보8개에도 미관측 필요성 등이 남아 있었다.
전체33,944명(18,467가구)을 유지하고 별도 요건 없는 일반 안내2개(가족전용상담전화, 마을변호사)만 별도 규칙 모델로 구현했다.
지원 여부를 검증한 학습모델은 만들 수 없으므로 실제 선정 예측 가능 정책수를2개라고 부르지 않는다.
국적 등 일부 집단으로 분석대상을 좁히면 기본 안내 후보가 달라질 수 있다. 이번2개는 전 표본·이전제외 유지라는 보수적인 범위이며 데이터의 이론적 최대 정책수가 아니다.

1인당 일반안내 수는 모든 응답자에서2개, 가중평균2.00개, 분산0이다. 이는 추천 개인화/성능 향상이 아니고 안내 대상의 보편성이다.
예측에 필요한 개인 특성0개. 만연령과 가구원가중값은 집계에만 사용하고 가구/가구원번호는 로컬 행 정렬에만 사용한다.
추가 질문, 결측 대치, 소득·등록자격 추정, 서로 다른 조사의 개인 단위 결합은 수행하지 않았다.
2025년 사회조사와 보관 정책 원문의 비교이며 현재 실제 이용/선정을 보장하지 않는다.

## 정책별 검토
'''
    for _,r in review.iterrows():report+=f"\n- {r['서비스명']}: {r['모델B_선정']} — {r['검토설명']}\n"
    report+='\n## 검증\n선정목록, 제외정책 누출, 가중집계, 개인정보 비공개, 원자료 해시를 별도로 검사한다. 학습/시험 분할과 정확도는 적용 대상이 아니다.\n'
    (out/'보고서.md').write_text(report,encoding='utf-8')
    dest=root/'발표자료/중간보고회발표자료/그림/무질문모델_20261003';dest.mkdir(parents=True,exist_ok=True)
    for p in list((out/'charts').glob('*.png'))+[out/'무질문모델_집계표.xlsx']:shutil.copy2(p,dest/p.name)
    return summary


def make_figures(root,out,groups):
    runtime=root/'.analysis_runtime'
    if runtime.exists():sys.path.append(str(runtime))
    os.environ.setdefault('MPLCONFIGDIR',str(runtime/'mplconfig'))
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False,'font.size':11})
    fig,ax=plt.subplots(figsize=(9,4.5),layout='constrained')
    ax.barh(['기존 활성 후보','미관측 조건 등으로 제외','일반 안내만 포함'],[8,6,2],color=['#7393A8','#C77942','#206B73'])
    ax.invert_yaxis();ax.set_xlim(0,9);ax.set_xlabel('서비스 수');ax.set_title('무질문 모델 선정: 8개 검토 → 일반 안내 2개')
    for i,v in enumerate([8,6,2]):ax.text(v+.12,i,str(v),va='center')
    fig.savefig(out/'charts/no_question_selection.png',dpi=170);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,4.8),layout='constrained')
    ax.bar(groups['생애주기'],groups['일반안내_가중평균'],color='#206B73',width=.55)
    ax.set_ylim(0,2.6);ax.set_ylabel('일반 안내 서비스 수 (가중평균)');ax.set_title('생애주기에 관계없이 일반 안내 2개')
    for i,r in groups.iterrows():ax.text(i,2.06,f"2.00\nn={int(r['표본수']):,}",ha='center',fontsize=10)
    fig.text(.5,-.045,'청년 19~34세 · 실제 지원금/수급 수가 아니며 개인화 성능을 의미하지 않음',ha='center',fontsize=9)
    fig.savefig(out/'charts/no_question_lifecycle.png',dpi=170,bbox_inches='tight');plt.close(fig)


def build_page(site,out,s,review,selected,groups,comparison,reasons):
    if not (site/'index.html').exists():raise ValueError('기존 사이트를 먼저 생성해야 합니다.')
    def table(t):return '<div class="table">'+t.to_html(index=False,escape=True,border=0,float_format=lambda x:f'{x:.2f}')+'</div>'
    body='''<h1>모델 B · 추가 질문 없는 분석</h1>
<p>질문하거나 개인 특성을 추정하지 않고, 현재 보유한 사회조사와 정책 원문만 사용했습니다. 기존 분석과 별도로 실행되는 보수적인 일반 안내 규칙 모델입니다.</p>
<div class="note"><strong>검토 결과의 한계</strong><br>실제 지원·선정 여부를 검증한 예측모델은 현재 자료로 만들지 못했습니다. 아래 2개는 별도 요건 없는 일반 정보·상담 안내이며, 수급 가능한 지원금 2개를 뜻하지 않습니다. 정확도나 실제 자격을 확인했다는 주장도 하지 않습니다.</div>
<div class="cards"><div class="card">검토한 기존 후보<b>8개</b>이전 제외 453개 유지</div><div class="card">일반 안내 규칙<b>2개</b>이번 모델에서 6개 추가 제외</div><div class="card">분석 표본<b>33,944명</b>18,467가구 · 전원 유지</div><div class="card">추가 질문·개인특성 추정<b>0개</b>학습하지 않는 상수 규칙</div></div>
<h2>1. 왜 더 엄격하게 선정했나</h2><p>연령·학력·복지카드가 관측되어도 상담 필요성, 디지털 어려움, 개별 프로그램 대상 여부까지 알 수 있는 것은 아닙니다. 미관측 조건을 충족했다고 가정하지 않고, 이 조건이 남는 서비스는 전체 표본용 모델에서 제외했습니다. 원문에 별도 지원요건이 없다고 적힌 일반 안내만 남겼습니다.</p>
<p>이전 제외 목록을 유지한 8개 후보를 다시 검토한 결과입니다. 461개를 이번에 새로 전수 재감사한 결과나 활용 가능한 정책의 이론적 최대치가 아닙니다. 특정 집단으로 범위를 제한하면 기본 안내 후보는 달라질 수 있습니다.</p>
<figure><img src="charts/no_question_selection.png" alt="8개 후보에서6개 제외하여2개 일반 안내 선정"></figure>
<h2>2. 새 모델이 출력하는 것</h2><p>각 응답자에게 선정된 서비스의 <strong>일반 안내 포함=1</strong>을 출력합니다. 두 서비스 모두 별도 요건이 없는 안내여서 전원 동일하게 2개가 나옵니다. 모델 밖 정책은 부적격=0으로 출력하지 않습니다. 실제 선정·지급 여부는 출력하지 않습니다.</p>
'''+table(selected[['서비스명','판정근거','일반안내_표본수','일반안내_가중비율','실제선정여부']])+'''
<p>가족상담의 심화 서비스·기관 연계, 마을변호사의 현장방문 가능 여부까지 포함하는 판정은 아닙니다. 이용 희망이나 필요를 관측했다고 간주하지 않으며 일반 안내 목록만 제공합니다.</p>
<h2>3. 생애주기별 결과</h2><p>가중평균 2.00개, 개인별 안내 수의 분산 0입니다. 보편적인 안내만 남겨 나타난 결과이며 높은 예측 성능이나 맞춤 추천을 의미하지 않습니다. 청년은 만19~34세입니다.</p>
'''+table(groups)+'''<figure><img src="charts/no_question_lifecycle.png" alt="모든 생애주기에서 일반 안내 가중평균2개"></figure>
<h2>4. 후보 8개의 포함·제외 근거</h2><p>정책별 검토설명과 보관 원문의 지원대상·선정기준을 함께 제공합니다. 제외는 이 사람에게 지원 불가라는 의미가 아닙니다.</p>
'''+table(review[['서비스명','모델B_선정','판정근거','검토설명']])+'''
<details><summary>검토에 사용한 정책 원문 보기</summary>'''+table(review[['서비스명','지원대상_보관원문','선정기준_보관원문','원문기준연도']])+'''</details>
<h2>5. 기존 분석과의 차이</h2>'''+table(comparison)+'''
<p>기존 평균6.02개와 새 평균2.00개는 서로 다른 범위를 측정합니다. 성능이 하락하거나 향상했다고 비교할 수 없습니다. 기존 보고서는 이용 희망 등을 전제한 기본대상 분석으로 유지합니다.</p>
<h2>6. 데이터·모델·평가 방법</h2><ul>
<li>사람 데이터: 사회조사2025, 33,944행·337열. 가구번호와 가구원번호 조합으로 중복을 확인했습니다.</li>
<li>정책 데이터: 보관 원문461개 및 기존 제외 목록. 새로운 설문, 외부 개인자료, 확산모델 합성자료는 사용하지 않았습니다. 가계동향조사를 같은 사람 자료처럼 연결하지 않았습니다.</li>
<li>실제 판정용 개인 특성은0개입니다. 만연령은 그룹 구분, 가구원가중값은 집계에만 사용합니다.</li>
<li>가중평균=Σ(개인 가중치×일반 안내 수)/Σ(개인 가중치). 외국 국적556명도 표본에 유지했습니다. 전표본 또는 각 그룹을 분모로 사용합니다.</li>
<li>실제 신청·승인·지급 정답은 없습니다. 모델이 만든 결과를 정답으로 재학습하지 않았습니다.</li>
<li>출력이 상수이므로 학습·시험 분할, 정확도, AUC, 확산모델 성능은 산출하지 않았습니다. 일반 안내100%는 추천 정확도100%가 아닙니다.</li>
<li>2025년 관측과 보관된 정책 문구의 비교입니다. 정책 기준연도 누락·2026년 문구가 있어 현재의 개인 자격이나 최신 운영 상황을 보장하지 않습니다.</li></ul>
<h2>7. 이전 감사에서 제외된451개: 대표 사유</h2><p>기존 감사 집계를 참고용으로 제공합니다. 정책당 대표 사유1개이며 한 정책의 모든 누락 조건을 세는 통계는 아닙니다. 여기에 기존19개 모델 정책의 추가 제외2개를 합쳐 이전 제외는453개입니다.</p>
'''+table(reasons)+'''
<h2>다운로드·재현</h2><p><a href="tables/no_question_무질문모델_집계표.xlsx">검토표·집계표 Excel</a> · <a href="tables/no_question_선정검토8개.csv">정책별 검토 CSV</a> · <a href="reproduce/no_question_model.json">규칙 모델 JSON</a> · <a href="reproduce/no_question_README.md">재현 방법</a> · <a href="reproduce/no_question_analysis.py">분석 코드</a></p>
<p class="small">개인별 판정 행과 원자료는 로컬에만 저장했습니다. 공개 자료는 정책 정보·집계·코드이며 개인기록을 포함하지 않습니다.</p>'''
    template=(site/'index.html').read_text(encoding='utf-8')
    template=re.sub(r'<title>.*?</title>','<title>추가 질문 없는 모델 B | 복지서비스 분석</title>',template,count=1)
    template=template.replace('2026-10-03 재분석 · 활성 서비스 8개 · 사회조사 2025','모델 B · 일반 안내 2개 · 사회조사 2025 · 추가 질문 없음')
    template=re.sub(r'<main>[\s\S]*?</main>',lambda m:'<main>'+body+'</main>',template,count=1)
    (site/'no_question_model.html').write_text(template,encoding='utf-8')
    marker='<!-- no-question-model-link -->'
    banner=marker+'<p class="note"><strong>별도 분석 추가:</strong> 추가 질문 없이 전 표본에서 판정하는 범위를 더 엄격하게 검토한 <a href="no_question_model.html">모델 B 보고서</a>를 확인하세요. 기존 8개는 이용 희망 등을 전제한 기본대상 분석입니다.</p><!-- /no-question-model-link -->'
    for p in site.glob('*.html'):
        text=p.read_text(encoding='utf-8')
        if 'href="no_question_model.html">무질문 모델 B</a>' not in text:
            text=text.replace('</nav>','<a href="no_question_model.html">무질문 모델 B</a></nav>',1)
        if p.name in ['index.html','mapping.html','variables.html','codebook.html','diffusion.html'] and marker not in text:
            text=text.replace('<main>','<main>'+banner,1)
        p.write_text(text,encoding='utf-8')


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
    args=parser.parse_args()
    print(json.dumps(build(args.root),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
