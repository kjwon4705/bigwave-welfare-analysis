# -*- coding: utf-8 -*-
"""제외 정책을 되살리지 않는 기본 이용대상 분석 및 공개 보고서 생성.

개인 원자료는 로컬에서만 읽는다. 공개 사이트에는 집계표/차트/방법만 저장한다.
판정 1=저장 원문의 기본대상 부합, 0=명시 기본조건 비해당, NaN=미확인.
실제 신청/선정/지급 예측이나 확산모델 학습이 아니다.
"""
from pathlib import Path
import hashlib
import html
import json
import re
import shutil
import sys
import os

_LOCAL_RUNTIME = Path(__file__).resolve().parents[1] / '.analysis_runtime'
if _LOCAL_RUNTIME.exists():
    sys.path.append(str(_LOCAL_RUNTIME))
os.environ.setdefault('MPLCONFIGDIR', str(_LOCAL_RUNTIME / 'mplconfig'))

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / '분석/Result/정책자격_정밀감사_20261003'
OUT = ROOT / '분석/Result/정책제외_재분석_20261003'
PROC = ROOT / '분석/Processed_data'
SITE = ROOT / '분석/Result/site'
CONFIG = ROOT / 'Code/활성정책_설정.json'
RAW_POLICY = ROOT / '분석/Raw_data/중앙부처복지서비스_데이터_보완.csv'
STAGES = ['청소년','청년','중장년','노년']
BLUE, ORANGE = '#286da8', '#df793a'
plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False,'font.size':10})

RULES = {
 '가족전용상담전화': ('all', '일반 가족을 포함하며 별도 선정기준 없음', '모든 응답자: 기본 안내'),
 '무료법률상담': ('national', '법률상담을 원하는 국민', '현재국적코드=1: 기본 안내; 그 외 미확인'),
 '마을변호사': ('all', '별도 지원요건 없이 모든 주민', '모든 응답자: 기본 안내'),
 '청소년성문화센터설치운영': ('national', '일반 국민 대상 성교육 서비스', '현재국적코드=1: 기본 안내; 그 외 미확인'),
 '노후준비서비스': ('national', '일반 국민 모두', '현재국적코드=1: 기본 안내; 그 외 미확인'),
 '디지털배움터': ('national', '디지털 교육을 원하는 국민', '현재국적코드=1: 기본 안내; 그 외 미확인'),
 '장애인 집합 정보화교육': ('card', '정보화교육을 원하는 장애인', '장애인복지카드_소유여부=1: 기본 안내; 미소유/결측은 미확인'),
 'WEE 클래스 상담지원': ('school', '초·중·고 재학생의 일반 상담', '교육정도코드∈{1,2,3} AND 재학상태코드=3; 결측 보존'),
}

def checked(path, base=ROOT):
    path,base=Path(path).resolve(),Path(base).resolve()
    if path==base or not path.is_relative_to(base): raise ValueError(f'범위 밖 경로: {path}')
    return path

def load_config():
    if CONFIG.exists(): return json.loads(CONFIG.read_text(encoding='utf-8'))
    a=pd.read_csv(AUDIT/'전체정책461.csv')
    d=pd.read_csv(AUDIT/'확산모델정책19.csv')
    bad=set(a.loc[a['감사분류']=='추가 확인 필요','서비스ID'])
    old=set(d['서비스ID'])
    active=set(a['서비스ID'])-(bad|old)
    assert len(bad)==451 and len(old)==19 and len(bad&old)==17 and len(active)==8
    c={'revision':'2026-10-03','original_count':461,'audit_excluded':sorted(bad),
       'retired_model_policies':sorted(old),'active_policy_ids':sorted(active),
       'source_sha256':hashlib.sha256(RAW_POLICY.read_bytes()).hexdigest(),
       'meaning':'기본 이용대상 안내; 실제 수급/선정 확정 아님; 이용희망 전제'}
    CONFIG.write_text(json.dumps(c,ensure_ascii=False,indent=2),encoding='utf-8')
    return c

def evaluate_basic(raw, policies):
    y={}
    for _,p in policies.iterrows():
        rule=p['판정규칙키']; n=len(raw)
        if rule=='all': arr=np.ones(n)
        elif rule=='national': arr=np.where(raw['현재국적코드'].eq(1),1.,np.nan)
        elif rule=='card': arr=np.where(raw['장애인복지카드_소유여부'].eq(1),1.,np.nan)
        elif rule=='school':
            edu=raw['교육정도코드'];enr=raw['재학상태코드']
            good=edu.isin([1,2,3]) & enr.eq(3)
            bad=(edu.notna() & ~edu.isin([1,2,3])) | (enr.notna() & ~enr.eq(3))
            arr=np.where(good,1.,np.where(bad,0.,np.nan))
        else: raise ValueError(f'검토되지 않은 규칙: {rule}')
        y[p['서비스ID']]=arr
    return pd.DataFrame(y,index=raw.index)

def csv(df,path): df.to_csv(path,index=False,encoding='utf-8-sig')

def grouped(raw,y,labels):
    w=raw['가구원가중값'];pos=y.eq(1).sum(axis=1);unk=y.isna().sum(axis=1)
    rows=[]
    for label in labels.dropna().unique():
        mask=labels.eq(label);ww=w[mask]
        rows.append({'구분':str(label),'표본수':int(mask.sum()),
         '기본대상부합_가중평균':float(np.average(pos[mask],weights=ww)),
         '추가확인_가중평균':float(np.average(unk[mask],weights=ww)),
         '명시비해당_가중평균':float(np.average(y.eq(0).sum(axis=1)[mask],weights=ww)),
         '중앙값_비가중':float(pos[mask].median())})
    return pd.DataFrame(rows)

def clean_axis(ax):
    ax.spines[['top','right']].set_visible(False)
    ax.grid(axis='x',alpha=.18);ax.set_axisbelow(True)

def figures(raw,y,policy_stats,stages,disability):
    charts=OUT/'charts';charts.mkdir(exist_ok=True)
    fig,ax=plt.subplots(figsize=(10.4,5.8),layout='constrained')
    p=policy_stats.iloc[::-1];left=np.zeros(len(p))
    for col,label,color in [('기본대상부합_가중%','기본대상 부합',BLUE),('추가확인_가중%','추가 확인','#bac4ce'),('명시비해당_가중%','명시 기본조건 비해당',ORANGE)]:
        ax.barh(p['서비스명'],p[col],left=left,label=label,color=color)
        left+=p[col].to_numpy()
    ax.set_xlim(0,100);ax.set_xlabel('전체 응답자 가중 비율 (%)')
    ax.set_title('남은 8개 서비스의 기본대상 분류',loc='left',pad=16)
    ax.legend(loc='lower center',bbox_to_anchor=(.5,-.28),ncol=3,frameon=False)
    fig.savefig(charts/'policy_status.png',dpi=170,bbox_inches='tight');plt.close(fig)
    age=raw['만연령']; stage=pd.cut(age,[12,18,34,64,np.inf],labels=STAGES)
    pos=y.eq(1).sum(axis=1);w=raw['가구원가중값']
    # 정수 0~8개이므로 밀도 평활화 대신 가중 이산분포를 그린다.
    dist=[]
    for g in STAGES:
        m=stage.eq(g)
        for k in range(9): dist.append({'생애주기':g,'기본대상부합수':k,'가중%':100*w[m&pos.eq(k)].sum()/w[m].sum()})
    dist=pd.DataFrame(dist);csv(dist,OUT/'생애주기별_분포.csv')
    fig,axes=plt.subplots(2,2,figsize=(10,7),sharex=True,sharey=True,layout='constrained')
    for ax,g,lab in zip(axes.flat,STAGES,['13~18세','19~34세','35~64세','65세 이상']):
        sub=dist[dist['생애주기']==g];ax.bar(sub['기본대상부합수'],sub['가중%'],color=ORANGE if g=='청년' else BLUE)
        ax.set_title(f'{g} ({lab})');ax.set_xticks(range(9));ax.set_ylim(0,100)
        ax.set_xlabel('기본대상 부합 서비스 수 (개)');ax.set_ylabel('그룹 내 가중 비율 (%)')
        ax.spines[['top','right']].set_visible(False)
    fig.suptitle('생애주기별 기본 이용대상 서비스 수 — 실제 수급 수 아님',fontsize=14)
    fig.savefig(charts/'lifecycle_distribution.png',dpi=170,bbox_inches='tight');plt.close(fig)
    for table,fname,title in [(stages,'lifecycle_mean.png','생애주기별 기본대상 부합 서비스 수'),(disability,'disability_mean.png','장애인복지카드 소유 응답별 기본대상 부합 수')]:
        fig,ax=plt.subplots(figsize=(8,4.5),layout='constrained');xx=table['기본대상부합_가중평균']
        bars=ax.barh(table['구분'],xx,color=BLUE)
        for bar,v in zip(bars,xx):ax.text(v+.08,bar.get_y()+bar.get_height()/2,f'{v:.2f}개',va='center')
        ax.set_xlim(0,8);ax.set_xlabel('가중 평균 (개, 전체 8개 중)');ax.set_title(title,loc='left');clean_axis(ax)
        fig.savefig(charts/fname,dpi=170,bbox_inches='tight');plt.close(fig)

def table(df):
    d=df.copy()
    for c in d:
        if pd.api.types.is_float_dtype(d[c]):d[c]=d[c].map(lambda v:f'{v:,.2f}' if pd.notna(v) else '미확인')
    return d.to_html(index=False,escape=True,border=0)

CSS='''*{box-sizing:border-box}body{margin:0;background:#f5f7fa;color:#203246;font:16px/1.75 'Malgun Gothic',sans-serif}header{background:#173b58;color:white;padding:34px max(20px,calc((100vw - 1060px)/2))}header h1{font-size:26px;margin:0}nav{background:white;border-bottom:1px solid #dce3ea;padding:14px 20px;display:flex;justify-content:center;gap:24px;flex-wrap:wrap}a{color:#236496}main{max-width:1100px;margin:auto;padding:25px 20px 65px}h1,h2{line-height:1.4}h2{margin-top:38px;font-size:22px}.note{background:#fff1d6;border-left:5px solid #bd7909;padding:18px;margin:22px 0}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:14px}.card{padding:22px;background:white;border:1px solid #dae2eb;border-radius:10px}.card b{display:block;font-size:34px;color:#1c537b}.table{overflow:auto}table{border-collapse:collapse;background:white;width:100%;font-size:14px}th,td{padding:11px 13px;border:1px solid #dae2eb;text-align:left;vertical-align:top;min-width:95px}th{background:#e9eff6}figure{margin:28px 0}img{max-width:100%;height:auto;border:1px solid #dce3ea;background:white}figcaption,.small{font-size:13px;color:#526779}footer{border-top:1px solid #dce3ea;text-align:center;font-size:13px;padding:22px}code{background:#e6edf4;padding:2px 5px}li{margin:6px 0}'''

def page(title,body):
    nav=''.join(f'<a href="{f}">{t}</a>' for f,t in [('index.html','홈'),('variables.html','중앙부처 복지서비스'),('eda.html','사회조사 EDA'),('mapping.html','공통변수'),('codebook.html','공통 변수 상세 설명'),('diffusion.html','분석방법 변경')])
    return f'<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{CSS}</style></head><body><header><h1>복지서비스 기본 이용대상 분석</h1><div>2026-10-03 재분석 · 활성 서비스 8개 · 사회조사 2025</div></header><nav>{nav}</nav><main>{body}</main><footer>복지로 보관 원문 · 통계청 사회조사 2025 · 원자료 비공개, 집계 결과만 공개</footer></body></html>'

def site_build(stats,policies,ps,st,ds,freqs):
    stage=OUT/'site_build';stage.mkdir(exist_ok=True)
    for folder in ['charts','tables','reproduce']:(stage/folder).mkdir(exist_ok=True)
    for file in (OUT/'charts').glob('*.png'):shutil.copy2(file,stage/'charts'/file.name)
    for file in ['정책별_기본대상.csv','생애주기별_평균.csv','생애주기별_분포.csv','장애카드별_평균.csv','제외집계.csv']:
        shutil.copy2(OUT/file,stage/'tables'/file)
    shutil.copy2(OUT/'재분석_집계표.xlsx',stage/'tables'/'analysis_tables.xlsx')
    # 공개 코드에는 마이크로데이터나 제외 정책 원문을 포함하지 않는다.
    shutil.copy2(Path(__file__),stage/'reproduce'/'policy_analysis.py')
    public_config=load_config()  # 공개된 정책의 ID 목록만 포함; 개인 식별정보 없음.
    (stage/'reproduce'/'active_config.json').write_text(json.dumps(public_config,ensure_ascii=False,indent=2),encoding='utf-8')
    (stage/'reproduce'/'README.md').write_text('''# 재현 방법\n개인 원자료는 공개하지 않습니다. pandas, numpy, matplotlib, openpyxl이 필요합니다.\n\n로컬 프로젝트에 다음 파일을 배치합니다.\n- policy_analysis.py → Code/정책제외_재분석.py\n- active_config.json → Code/활성정책_설정.json\n- 보관 정책 원문 → 분석/Raw_data/중앙부처복지서비스_데이터_보완.csv\n- 이용 허가를 받은 사회조사 원자료 → 분석/Raw_data/2025_*.csv (한 파일)\n\n분석/Processed_data 및 발표자료/중간보고회발표자료/그림 폴더를 만든 뒤 프로젝트 루트에서 Python으로 Code/정책제외_재분석.py를 실행합니다.\n정책 원문 해시가 보관본과 다르면 중단합니다. 개인 원자료가 없는 공개 저장소만으로는 다시 계산할 수 없습니다.\n가중치는 가구원가중값을 사용하고 unknown은 비율 분모에 포함합니다. 공개 설정은 정책ID 목록이며 개인ID가 아닙니다.\n''',encoding='utf-8')
    note='<p class="note">여기서 기본대상 부합은 저장된 정책 원문의 일반 이용대상 조건과 일치한다는 뜻입니다. 서비스 이용 희망을 전제로 하며, 실제 신청 승인·선정·급여 지급을 예측하지 않습니다. 미확인을 부적격으로 처리하지 않았습니다.</p>'
    cards=f'<div class="cards"><div class="card">활성 서비스<b>8개</b></div><div class="card">분석 대상<b>33,944명</b>18,467가구</div><div class="card">기본대상 부합 평균<b>{stats["weighted_mean_positive"]:.2f}개</b>가구원 가중치 적용</div><div class="card">추가확인 평균<b>{stats["weighted_mean_unknown"]:.2f}개</b>개인별 미확인 조건</div></div>'
    flow='<p>원정책 461개에서 추가확인 대상 451개와 기존 확산모델의 19개를 모두 제외했습니다. 두 목록에서 17개가 겹쳐 총 453개를 제외했고, 8개를 재분석했습니다.</p>'
    downloads='<p><a href="tables/analysis_tables.xlsx">집계표 엑셀 내려받기</a> · <a href="tables/정책별_기본대상.csv">정책별 집계 CSV</a> · <a href="reproduce/README.md">재현 방법과 코드</a></p>'
    home='<h1>제외 정책 반영 후 재분석</h1>'+flow+cards+note+'<h2>이번 결과가 말하는 것</h2><p>일반 안내 서비스가 다수여서 평균값이 높게 나타납니다. 이를 지원금의 풍부함이나 개인별 실제 수혜량으로 해석할 수 없습니다. 개인별 미확인 항목은 별도로 남겼습니다.</p><figure><img src="charts/policy_status.png" alt="활성 8개 서비스별 기본대상 부합, 추가확인, 명시비해당 가중 비율"></figure>'+downloads
    variables='<h1>중앙부처 복지서비스</h1>'+flow+note+'<h2>현재 분석하는 8개 서비스</h2><div class="table">'+table(policies[['서비스명','소관부처','기본대상_원문요약','사람쪽_확인방법']])+'</div><h2>기본대상 분류 결과</h2><div class="table">'+table(ps)+'</div><p>비율 분모는 정책마다 동일한 전체 응답자 가중치 합입니다. 추가확인 응답자를 분모에서 빼지 않습니다.</p>'+downloads
    mapping='<h1>공통변수와 재분석 결과</h1>'+note+'<p>현재 판정에는 국적·복지카드 소유·교육정도·재학상태만 사용합니다. 연령은 생애주기별 집계를 위한 변수입니다. 소득·수급자 근사지표는 자격 조건으로 사용하지 않습니다.</p><h2>생애주기별 평균</h2><div class="table">'+table(st)+'</div><figure><img src="charts/lifecycle_mean.png" alt="생애주기별 기본대상 부합 가중 평균"></figure><h2>생애주기별 분포</h2><figure><img src="charts/lifecycle_distribution.png" alt="생애주기별 정수 서비스 수 분포"><figcaption>청년은 만19~34세입니다. 서비스 수가 0~8의 정수이므로 연속 밀도처럼 보일 수 있는 바이올린 대신 가중 빈도 막대를 사용했습니다.</figcaption></figure><h2>장애인복지카드 소유 응답별</h2><div class="table">'+table(ds)+'</div><figure><img src="charts/disability_mean.png" alt="복지카드 응답별 기본대상 부합 가중 평균"></figure><p>카드 미소유를 미등록장애로 단정하지 않습니다. 장애인 정보화교육에서는 카드 소유 응답에 기본 안내를 부여하고 나머지는 추가확인으로 남겼습니다.</p><h2>수급자별 비교</h2><p>사회조사에서 법정 기초생활수급·차상위 자격이 직접 관측되지 않으므로 해당 비교를 산출하지 않았습니다. 생활비 마련방법을 수급자 여부로 바꾸지 않습니다.</p>'+downloads
    method='<h1>기존 확산모델 분석 철회와 현재 분석방법</h1><p class="note">기존에 사용한 19개 정책을 이번 분석에서 모두 제외했습니다. 그 정책들의 예측 결과·성능표·합성 비교 그림은 현재 보고서에서 제공하지 않습니다.</p>'+flow+'<h2>현재 수행한 분석</h2><p>사회조사 원변수에 명시적인 기본대상 규칙을 적용하고, 전체 표본을 분모로 부합·비해당·미확인을 집계했습니다. 학습이나 개인 특성 추정은 수행하지 않았습니다. 따라서 이 보고서에는 새로운 확산모델 성능이나 예측 정확도 수치가 없습니다.</p><h2>모델을 다시 학습하지 않은 이유</h2><p>실제 신청·선정 결과라는 정답이 없고 남은 서비스는 대부분 일반 상담·교육입니다. 이전의 불완전한 자격규칙을 새 모델로 학습해도 실제 자격 검증을 대신할 수 없습니다. 향후 실제 결과 또는 검증된 정책별 신청경로를 확보한 뒤 별도 평가해야 합니다.</p><h2>해석 범위</h2><p>2025년 조사 속성과 보관된 정책 조건을 비교한 분석입니다. 현재 개인의 정보, 최신 공고, 접수기간을 확인한 결과가 아닙니다. 표본설계를 반영한 가중 비율은 제시하지만 설계기반 신뢰구간이나 유의성 검정은 산출하지 않았습니다.</p>'
    cb=pd.DataFrame([
      ('판정 결과','1 / 0 / 미확인','1=기본조건 부합, 0=명시조건 비해당, 미확인=자료 부족. 행정 자격 아님'),
      ('국적','현재국적코드','1 대한민국이면 국민 대상 일반 안내. 외국국적/결측은 일괄 탈락시키지 않음'),
      ('등록장애의 대리관측','장애인복지카드_소유여부','소유1은 장애인 교육 안내에 활용. 미소유2와 결측은 등록상태 미확인'),
      ('학교 단계','교육정도코드','1 초등학교, 2 중학교, 3 고등학교. 기존 학생 하나의 변수로 대체하지 않음'),
      ('재학','재학상태코드','3 재학. 4 휴학을 자동으로 재학과 동일시하지 않음. 결측 유지'),
      ('WEE 일반상담','학교 단계 AND 재학','두 조건 모두 참이면1; 관측된 필수조건 하나라도 거짓이면0; 나머지는 미확인'),
      ('생애주기','만연령','13~18 청소년, 19~34 청년, 35~64 중장년, 65+ 노년. 집계용이며 정책연령 제한 아님'),
      ('가중 집계','가구원가중값','전 표본 또는 해당 그룹의 가중치 합을 고정분모로 사용. 미확인을 분모에서 제거하지 않음'),
      ('서비스 수','부합 서비스 수의 합','부합=1의 개수만 합산. 미확인 수를 별도로 함께 제시. 실제 수급 수 아님'),
    ],columns=['항목','사용 원변수/규칙','생성·해석 기준'])
    codebook='<h1>공통 변수 상세 설명</h1>'+note+'<div class="table">'+table(cb)+'</div><h2>정책별 연결</h2><div class="table">'+table(policies[['서비스명','기본대상_원문요약','사람쪽_확인방법']])+'</div><h2>제외 원칙</h2><p>활성 정책 ID 목록에 포함된 8개만 계산합니다. 기존 모델의 19개 ID는 금지 목록으로 관리하여 재실행해도 다시 포함되지 않도록 검사합니다. 과거의 태그 기반 0/1/2 조건이나 소득·수급 근사값은 새 판정에 사용하지 않습니다.</p>'
    intro="본 분석은 총 18,467가구에 거주하는 만 13세 이상 가구원 33,944명을 최종 분석 대상으로 설정하였다. 분석 결과 제시 시 원자료 표본 집계 결과인 '표본 비율(%)'과 함께, 전국 모집단 대표성을 확보하도록 설계된 복합표본 가중치를 적용한 '가중 비율(%)'을 병기하였다. 주요 변수 및 응답 항목의 명칭은 통계청이 제공하는 원자료 코드북의 표준 분류 기준을 따랐다."
    eda='<h1>사회조사 2025 원변수 기초 EDA</h1><p>'+intro+'</p><p>현재 분석에 사용하는 원변수와 생애주기 분포를 제시합니다. 비해당·결측도 별도 범주로 포함하여 전체 표본 기준 비율을 계산했습니다.</p>'
    for name,df in freqs.items():eda+='<h2>'+html.escape(name)+'</h2><div class="table">'+table(df)+'</div>'
    # 정책 제외와 무관한 원변수 29개 EDA는 유지한다. 정책 그림은 복사하지 않는다.
    oldsite=ROOT/'분석/Archive/20261003_재분석이전/분석/Result/site'
    if (oldsite/'eda.html').exists():
        oldhtml=(oldsite/'eda.html').read_text(encoding='utf-8')
        match=re.search(r'<main>([\s\S]*?)</main>',oldhtml)
        if match:
            eda=match.group(1)
            eda=re.sub(r'<div><b>[^<]*</b><span>자녀 있음</span></div>','',eda)
            eda=eda.replace('<h1>사회조사 2025 원변수 기초 EDA</h1>','<h1>사회조사 2025 원변수 기초 EDA</h1><p class="note">원변수 EDA는 정책 제외와 독립적이므로 유지했습니다. 이 페이지의 범주별 비율은 비결측 응답자를 분모로 합니다. 자녀 관련 문항은 전체 응답자의 실제 자녀 유무가 아니며, 연령 평균은 비가중 표본 평균입니다.</p>')
            for img in (oldsite/'charts').glob('chart_*.png'):
                if re.fullmatch(r'chart_\d{2}\.png',img.name):shutil.copy2(img,stage/'charts'/img.name)
    eda=eda.replace('<table','<table')
    pages={'index.html':home,'variables.html':variables,'mapping.html':mapping,'diffusion.html':method,'codebook.html':codebook,'eda.html':eda}
    for name,body in pages.items():(stage/name).write_text(page(name,body),encoding='utf-8')
    (stage/'.nojekyll').touch()
    (stage/'README.md').write_text('# 복지서비스 기본 이용대상 분석\n\n2026-10-03: 461개 중 453개를 분석에서 제외하고 8개 일반 서비스의 기본 이용대상을 재분석했습니다. 이전 확산모델 결과는 철회했습니다.\n\n개인 원자료와 행정 신청/수급 정답은 공개하지 않습니다. 공개 파일은 집계 결과·차트·설명·재현용 코드입니다.\n',encoding='utf-8')
    (stage/'analysis_summary.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    # 검증된 신규 공개 산출물만 원래 사이트 경로에 동기화; 과거 그래프의 직접 URL도 제거.
    expected={p.relative_to(stage) for p in stage.rglob('*') if p.is_file()}
    SITE.mkdir(exist_ok=True)
    for old in SITE.rglob('*'):
        if old.is_file() and old.relative_to(SITE) not in expected:checked(old,SITE).unlink()
    shutil.copytree(stage,SITE,dirs_exist_ok=True)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    c=load_config()
    assert hashlib.sha256(RAW_POLICY.read_bytes()).hexdigest()==c['source_sha256'],'정책 원문 변경: 먼저 재감사 필요'
    rawpol=pd.read_csv(RAW_POLICY).fillna('')
    active=set(c['active_policy_ids']);removed=set(c['audit_excluded'])|set(c['retired_model_policies'])
    assert len(active)==8 and len(removed)==453 and not active&removed
    sel=rawpol[rawpol['서비스ID'].isin(active)].copy()
    assert set(sel['서비스명'])==set(RULES) and sel['서비스ID'].is_unique
    pol=sel[['서비스ID','서비스명','소관부처','복지로 상세링크']].copy()
    pol['판정규칙키']=pol['서비스명'].map(lambda n:RULES[n][0])
    pol['기본대상_원문요약']=pol['서비스명'].map(lambda n:RULES[n][1])
    pol['사람쪽_확인방법']=pol['서비스명'].map(lambda n:RULES[n][2])
    pol['분석의미']='이용 희망 전제의 기본 안내; 실제 자격 확정 아님'
    rawfile=next((ROOT/'분석/Raw_data').glob('2025_*.csv'))
    raw=pd.read_csv(rawfile,encoding='cp949',low_memory=False)
    keys=['가구일련번호','가구원번호'];assert not raw.duplicated(keys).any()
    w=raw['가구원가중값'];assert w.notna().all() and (w>0).all()
    y=evaluate_basic(raw,pol);assert set(y)==active
    assert y.isin([0,1]).to_numpy().sum()+y.isna().to_numpy().sum()==len(raw)*8
    result=pd.concat([raw[keys],y],axis=1)
    csv(pol,PROC/'중앙부처복지서비스_변수화.csv');csv(pol,PROC/'중앙부처복지서비스_활성정책.csv')
    csv(sel,PROC/'중앙부처복지서비스_활성원문.csv')
    csv(result,PROC/'매칭결과_사회조사.csv')
    csv(raw[keys+['가구원가중값','만연령','현재국적코드','장애인복지카드_소유여부','교육정도코드','재학상태코드']],PROC/'사회조사_활성원변수.csv')
    excluded=rawpol.loc[rawpol['서비스ID'].isin(removed),['서비스ID','서비스명']].copy()
    excluded['감사추가확인']=excluded['서비스ID'].isin(c['audit_excluded'])
    excluded['기존19개전부제외']=excluded['서비스ID'].isin(c['retired_model_policies'])
    csv(excluded,PROC/'제외정책목록.csv')
    csv(pol[['서비스ID','서비스명']].assign(검토상태='기본안내만; 행정자격 미검증'),PROC/'정책Tier목록.csv')
    olddir=PROC/'확산모델';olddir.mkdir(exist_ok=True)
    csv(pd.DataFrame(columns=['서비스ID','서비스명','선정']),olddir/'정책선정.csv')
    ps=[]
    for _,row in pol.iterrows():
        v=y[row['서비스ID']];r={'서비스ID':row['서비스ID'],'서비스명':row['서비스명']}
        for mask,label in [(v.eq(1),'기본대상부합'),(v.eq(0),'명시비해당'),(v.isna(),'추가확인')]:
            r[label+'_표본수']=int(mask.sum());r[label+'_가중%']=100*w[mask].sum()/w.sum()
        ps.append(r)
    ps=pd.DataFrame(ps)
    ages=pd.cut(raw['만연령'],[12,18,34,64,np.inf],labels=STAGES)
    st=grouped(raw,y,ages).set_index('구분').reindex(STAGES).reset_index()
    dl=raw['장애인복지카드_소유여부'].map({1:'카드 소유',2:'카드 미소유'}).fillna('카드 응답 미확인')
    ds=grouped(raw,y,dl)
    stats={'revision':'2026-10-03','analysis_type':'기본 이용대상 규칙 집계; 확산모델 아님',
      'original_policies':461,'audit_excluded':451,'retired_model_policies':19,'overlap':17,'excluded_union':453,'active_policies':8,
      'people':len(raw),'households':raw['가구일련번호'].nunique(),
      'weighted_mean_positive':float(np.average(y.eq(1).sum(axis=1),weights=w)),
      'weighted_mean_unknown':float(np.average(y.isna().sum(axis=1),weights=w)),
      'weighted_mean_negative':float(np.average(y.eq(0).sum(axis=1),weights=w)),
      'active_service_names':pol['서비스명'].tolist(),'raw_policy_sha256':c['source_sha256'],
      'raw_person_sha256':hashlib.sha256(rawfile.read_bytes()).hexdigest(),
      'denominator':'전 표본/각 집계그룹 가구원 가중치 합; 미확인 포함'}
    assert np.isclose(stats['weighted_mean_positive']+stats['weighted_mean_unknown']+stats['weighted_mean_negative'],8)
    flow=pd.DataFrame({'단계':['원정책','감사 추가확인','기존 모델 전부','두 제외목록 중복','총 제외','최종 활성'],'정책수':[461,451,19,17,453,8]})
    for df,name in [(ps,'정책별_기본대상.csv'),(st,'생애주기별_평균.csv'),(ds,'장애카드별_평균.csv'),(flow,'제외집계.csv')]:csv(df,OUT/name)
    freqs={}
    fs={'생애주기':ages.astype('object'),'현재국적':raw['현재국적코드'].map({1:'대한민국',2:'외국'}),
       '장애인복지카드':dl,'교육정도':raw['교육정도코드'].map({0:'받지 않음',1:'초등학교',2:'중학교',3:'고등학교',4:'대학(4년제 미만)',5:'대학교(4년제 이상)',6:'대학원 석사',7:'대학원 박사'}),
       '재학상태':raw['재학상태코드'].map({1:'졸업',2:'수료',3:'재학',4:'휴학',5:'중퇴'})}
    for title,series in fs.items():
        v=series.fillna('비해당·결측');t=pd.DataFrame({'범주':v,'w':w}).groupby('범주',sort=False).agg(표본수=('w','size'),가중합=('w','sum')).reset_index()
        t['표본%']=100*t['표본수']/len(raw);t['가중%']=100*t['가중합']/w.sum();t=t.drop(columns='가중합');freqs[title]=t
        csv(t,OUT/f'EDA_{title}.csv')
    figures(raw,y,ps,st,ds)
    with pd.ExcelWriter(OUT/'재분석_집계표.xlsx',engine='openpyxl') as writer:
        for name,df in [('활성8개',pol),('제외집계',flow),('정책별집계',ps),('생애주기별',st),('장애카드별',ds),*freqs.items()]:
            df.to_excel(writer,sheet_name=name,index=False)
            ws=writer.book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=40 if col[0].value in ['서비스명','기본대상_원문요약','사람쪽_확인방법'] else 23
            from openpyxl.styles import Alignment, Font, PatternFill
            for cell in ws[1]:
                cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244F73');cell.alignment=Alignment(wrap_text=True)
            ws.row_dimensions[1].height=32
            for row in ws.iter_rows(min_row=2):
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
                ws.row_dimensions[row[0].row].height=46
    (OUT/'summary.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    report=f'''# 제외 정책 반영 후 재분석\n\n원정책 461개 중 감사 추가확인451개와 기존 모델19개를 모두 제외했습니다. 중복17개이므로 제외합집합453개, 활성8개입니다.\n\n분석대상은 {len(raw):,}명 / {stats['households']:,}가구입니다. 기본대상 부합 가중평균은 {stats['weighted_mean_positive']:.3f}개, 추가확인은 {stats['weighted_mean_unknown']:.3f}개입니다. 실제 수급정책 수가 아닙니다.\n\n남은 서비스는 일반 상담·교육 위주입니다. 이용 의사를 전제한 명시적 규칙 분석이며 모델 학습을 하지 않았습니다. 기존19개 확산모델 성능·예측 결과는 철회합니다. 수급자별 비교도 직접 관측된 자격이 없어 산출하지 않습니다.\n\n활성 서비스: {', '.join(pol['서비스명'])}.\n\n변경된 데이터는 Processed_data의 활성정책/매칭결과에 저장했습니다. 원본461개는 추적용으로 Raw_data에 보존했습니다. 이전 결과·코드·모델은 비공개 Archive에 보관합니다. 공개 사이트에는 개인기록을 포함하지 않습니다.\n'''
    (OUT/'재분석보고서.md').write_text(report,encoding='utf-8')
    site_build(stats,pol,ps,st,ds,freqs)
    # 별도 모델 B 페이지도 함께 재생성하여 기존 보고서 갱신 때 유실되지 않게 한다.
    from 무질문모델 import build as build_no_question_model
    build_no_question_model(ROOT)
    # 발표용 최신 그림은 별도 폴더에 동일 산출물을 복사.
    dest=ROOT/'발표자료/중간보고회발표자료/그림/재분석_20261003';dest.mkdir(exist_ok=True)
    for f in list((OUT/'charts').glob('*.png'))+list(OUT.glob('*.csv'))+[OUT/'재분석_집계표.xlsx']:shutil.copy2(f,dest/f.name)
    print(json.dumps(stats,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
