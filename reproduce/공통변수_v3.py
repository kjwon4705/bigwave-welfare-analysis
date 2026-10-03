"""검토한 조건만 구조화하고 미검토 정책을 잠그는 재현 가능한 전처리 v3."""
from pathlib import Path
import hashlib
import json
import shutil
import numpy as np
import pandas as pd
from 조건판정_v3 import VERSION,make_facts,curated_rules,evaluate,atom,all_of,any_of,unknown

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'분석/Result/공통변수_수정결과_20261003'
PROC=ROOT/'분석/Processed_data'
DEST=PROC/'공통변수_v3'
CONFIG=ROOT/'Code/활성정책_설정.json'


def rules_for(policies):
    rules=curated_rules()
    byname=policies.set_index('서비스명')['서비스ID']
    rules[byname['직장인 든든한 점심밥']]=(
        all_of(atom('임금근로_관측','eq',1),unknown('중소기업기본법상 중소기업 재직'),unknown('기업의 식대 지급 및 해당 사업 참여')),
        '부분구조화','취업 전체가 아닌 임금근로자 조건. 사업장 규모·식대 지급은 추가 요건.')
    rules[byname['체불 근로자 생계비 융자']]=(
        all_of(any_of(all_of(atom('임금근로_관측','eq',1),unknown('폐업하지 않은 체불사업장 재직')),
                      unknown('6개월 이내 퇴직 경로'),unknown('건설일용근로자 근로일수 경로')),
               unknown('체불기간·금액·신용 제외요건')),
        '부분구조화','재직자OR퇴직자OR건설일용 특례. 미취업만으로 정책 전체를 제외하지 않음.')
    if not set(rules)<=set(policies['서비스ID']):raise ValueError('검토 정책ID 불일치')
    return rules


def flatten(expr,path='root'):
    if 'all' in expr or 'any' in expr:
        kind='all' if 'all' in expr else 'any'
        for i,child in enumerate(expr[kind]):yield from flatten(child,f'{path}/{kind}[{i}]')
    elif 'not' in expr:yield from flatten(expr['not'],path+'/not')
    else:yield path,expr


def schema_for(policies,config):
    rules=rules_for(policies);records=[]
    excluded=set(config['audit_excluded'])|set(config['retired_model_policies'])
    for _,p in policies.iterrows():
        pid=p['서비스ID']
        expr,state,note=rules.get(pid,(unknown('전체 조건 구조화·원문 검토 미완료'),'미검토_잠금','태그/키워드에서 자격조건을 자동 생성하지 않음'))
        records.append({'서비스ID':pid,'서비스명':p['서비스명'],'검토상태':state,'조건식':expr,
          '검토설명':note,'기준연도':p['기준연도'],'대상주체':'서비스 이용 본인; 행정 신청대리 관계는 미검증',
          '출처':p['복지로 상세링크'],'지원대상_원문':p['지원대상 상세'],'선정기준_원문':p['선정기준'],
          '지원내용_원문':p['지원내용 상세'],'전체자격검증완료':False,
          '현재추천목록포함':pid in set(config['active_policy_ids']), '기존제외유지':pid in excluded,
          '출력의미':'관측조건 판정 또는 일반안내. 실제선정/지급 정답이 아님.'})
    return records


def load():
    c=json.loads(CONFIG.read_text(encoding='utf-8'))
    pp=ROOT/'분석/Raw_data/중앙부처복지서비스_데이터_보완.csv'
    if hashlib.sha256(pp.read_bytes()).hexdigest()!=c['source_sha256']:raise ValueError('정책 원문 변경: 재검토 필요')
    files=list((ROOT/'분석/Raw_data').glob('2025_*.csv'))
    if len(files)!=1:raise ValueError('사회조사 입력을 하나로 특정해야 합니다.')
    return c,pd.read_csv(pp).fillna(''),pd.read_csv(files[0],encoding='cp949',low_memory=False),files[0]


def active_evaluate(raw,policies):
    """현재8개 보고서도 동일한 v3 조건식 사용. 기존 제외정책은 재포함하지 않음."""
    c,source,_,_=load()
    if not set(policies['서비스ID'])<=set(c['active_policy_ids']):raise ValueError('활성목록 밖 정책')
    rules=rules_for(source);facts=make_facts(raw)
    return pd.DataFrame({pid:evaluate(rules[pid][0],facts) for pid in policies['서비스ID']},index=raw.index)


def build():
    for p in [OUT,DEST]:p.mkdir(parents=True,exist_ok=True)
    c,policies,raw,rawfile=load();facts=make_facts(raw);records=schema_for(policies,c)
    facts.to_csv(DEST/'사회조사_관측사실.csv',index=False,encoding='utf-8-sig')
    # 정식 표준변수 파일도 새 스키마로 교체. 구47열 원본은 수정전 Archive에 보존.
    facts.to_csv(PROC/'사회조사_표준변수.csv',index=False,encoding='utf-8-sig')
    (DEST/'정책조건.json').write_text(json.dumps({'version':VERSION,'policy_source_sha256':c['source_sha256'],'policies':records},ensure_ascii=False,indent=2),encoding='utf-8')
    rows=[];ps=[]
    w=facts['가구원가중값']
    for rec in records:
        y=evaluate(rec['조건식'],facts)
        for path,node in flatten(rec['조건식']):rows.append({'서비스ID':rec['서비스ID'],'서비스명':rec['서비스명'],'논리경로':path,
           '주체':node.get('subject','미확인/상수'),'속성':node.get('field',''),'연산':node.get('op',''),
           '값':json.dumps(node.get('value'),ensure_ascii=False),'조건역할':node.get('role','미확인/일반안내'),
           '미확인사유':node.get('unknown',''),'상수안내':node.get('constant',''),'검토상태':rec['검토상태'],
           '출처':rec['출처'],'근거필드':'지원대상 상세 + 선정기준','검토설명':rec['검토설명']})
        ps.append({'서비스ID':rec['서비스ID'],'서비스명':rec['서비스명'],'검토상태':rec['검토상태'],
          '기존제외유지':rec['기존제외유지'],'현재추천목록포함':rec['현재추천목록포함'],
          '관측조건부합_표본수':int((y==1).sum()),'명시조건비해당_표본수':int((y==0).sum()),'미확인_표본수':int(np.isnan(y).sum()),
          '부합_가중비율':float(100*w[y==1].sum()/w.sum()),'미확인_가중비율':float(100*w[np.isnan(y)].sum()/w.sum()),
          '해석':'미검토/부분구조화 결과는 진단용. 추천목록 확대나 실제자격 판정으로 사용 금지.'})
    table=pd.DataFrame(rows);stats=pd.DataFrame(ps)
    table.to_csv(DEST/'정책조건_장형.csv',index=False,encoding='utf-8-sig')
    stats.to_csv(OUT/'전체461_수정후진단.csv',index=False,encoding='utf-8-sig')
    dictionary=pd.DataFrame([{'변수명':col,'뜻':'미관측 법정/행정자격' if col.endswith('_법정') or col in ['소득인정액','질병진단'] else '원관측 또는 명시적 관측사실; 이름과 코드북 참조',
        '결측수':int(facts[col].isna().sum()),'자격추정허용':False if col.endswith('_법정') else '일반 사실만; 개별 정책 정의 대조 필요'} for col in facts])
    dictionary.to_csv(DEST/'사람변수사전.csv',index=False,encoding='utf-8-sig')
    fixes=pd.DataFrame([
       ['F01','태그는 검색정보로만 유지','태그에서 연령/전용 여부 생성 제거'],
       ['F02','연령 경로·예외 분리','장애수당 상한20세 삭제; 아동수당 경로OR; 월말 생일경계 미확인'],
       ['F03','부정문 키워드 추출 제거','학생/자영업 제외를 필수로 생성하지 않음'],
       ['F04','조건마다 대상주체 저장','본인·자녀를 임의대체하지 않음'],
       ['F05','정책별 AND/OR/NOT 식','경로OR·경로내AND·퇴직경로 보존'],
       ['F06','급여종류·수급권 분리','소득% 병기로 수급 필수조건을 완화하지 않음'],
       ['F07','생활비 원뜻 보존','수급/차상위는미관측'],
       ['F08','소득/가구수 구간 보존','4인이상 상한미공개;800만원이상 상한없음;인정액미추정'],
       ['F09','동읍면3/4 명시매핑','법정 농어촌·도서벽지로 확대해석하지 않음'],
       ['F10','미검토 정책 잠금','조건추출실패=미확인;무조건통과 삭제'],
       ['F11','관측과 법정자격 분리','카드/국적/임차에서 법정상태를 추정하지 않음'],
       ['F12','조사대상·결측 보존','60세미만 자녀문항미관측;무학 재학상태 비해당 구분'],
       ['F13','완전자격검증 상태 분리','NaN없음으로TierA생성 금지'],
       ['F14','전체표본분모 유지','미확인 포함한 부합/비해당/미확인 집계'],
       ['F15','현행8개 규칙도 v3연결','상담/디지털 어려움 누락은미확인 처리'],
       ['F16','2개 일반안내 의미 유지','실제 수급예측이라고 표현하지 않음'],
       ['F17','레거시 실행경로 교체/차단','새 전처리 연결, 미수정 구 스크립트는실행차단'],
       ['F18','사실·조건·결과 스키마 분리','공통0/1/2 숫자 직접비교 제거'],
    ],columns=['감사번호','수정','처리내용'])
    fixes.to_csv(OUT/'수정내역18개.csv',index=False,encoding='utf-8-sig')
    state=stats['검토상태'].value_counts().to_dict()
    summary={'version':VERSION,'people':len(raw),'households':int(raw['가구일련번호'].nunique()),'policies':len(records),
       'reviewed_policy_structures':int(sum(v for k,v in state.items() if k!='미검토_잠금')),
       'unreviewed_locked':int(state.get('미검토_잠금',0)),'state_counts':state,
       'fully_validated_eligibility':0,'prior_exclusions_preserved':453,
       'source_people_sha256':hashlib.sha256(rawfile.read_bytes()).hexdigest(),
       'source_policy_sha256':c['source_sha256'],'meaning':'구현 오류 수정 및 미검토 정책 잠금. 전체 정책 조건을 수동 완전검증한 것은 아님.'}
    (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    with pd.ExcelWriter(OUT/'공통변수_수정결과.xlsx',engine='openpyxl') as writer:
        for name,df in [('수정내역',fixes),('정책진단461',stats),('검토조건',table[table['검토상태'].ne('미검토_잠금')]),('사람변수사전',dictionary)]:
            df.to_excel(writer,sheet_name=name,index=False);ws=writer.book[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            from openpyxl.styles import Alignment,Font,PatternFill
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244F73')
            for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=38
            for row in ws:
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    (OUT/'수정보고서.md').write_text(f'''# 공통변수 v3 수정 결과

감사에서 확인한18개 문제 유형에 대응하는 구현을 변경했다. 실제 자격정답을 확보하거나 모든461개 정책을 완전히 재작성했다는 뜻은 아니다.

기존 원자료는 보존했고 잘못된47열 사회조사 표준변수는 관측사실 스키마로 교체했다. 과거코드와 표는 분석/Archive/20261003_공통변수수정전에 보존했다.
정책은 AND/OR/NOT,주체,원문,검토상태가 있는 구조로 바꿨다. 대표12개 및 기존8개 후보의 합집합인{summary['reviewed_policy_structures']}개를 구조화했다.
나머지{summary['unreviewed_locked']}개는 미검토 잠금이며 관측조건을 임의추정하지 않는다. 기존 제외453개는 추천에서 계속 제외한다.
모든 정책의 전체자격검증완료 값은 false이며 부분구조화는 실제 선정 보장이 아니다.

현재8개 보고서도 수정 조건을 적용하므로 기존6.02개 평균을 유지하지 않는다. 일반안내2개 모델B의 의미는 유지한다.
수정 전후 반례,3값논리,결측/구간,원자료보존 및 제외정책 재포함 여부를 회귀 검사한다.
''',encoding='utf-8')
    return summary


def append_site():
    """분석 페이지 생성 후 수정 설명/코드를 붙이고 오래된 의미 설명을 바로잡는다."""
    import re
    site=ROOT/'분석/Result/site'
    s=json.loads((OUT/'summary.json').read_text(encoding='utf-8'))
    current=json.loads((site/'analysis_summary.json').read_text(encoding='utf-8'))
    fixes=pd.read_csv(OUT/'수정내역18개.csv')
    body=f'''<h1>공통변수 v3 · 오류 수정 결과</h1>
<p>정책 원문의 필수·제외·예외조건을 구분하고 사람의 관측사실과 행정상 자격을 분리했습니다. 자동 추출이 실패한 정책을 무조건 통과시키는 동작을 제거했습니다.</p>
<div class="note"><strong>수정 범위와 남은 검토</strong><br>대표12개와 기존8개 후보의 합집합16개 조건을 구조화했습니다. 그중13개는 일부 조건을 구조화한 상태입니다. 나머지445개는 미검토 잠금입니다. 실제 선정·지급 정답으로 검증 완료한 정책은 아직 없습니다.</div>
<div class="cards"><div class="card">정책 원문<b>461개</b></div><div class="card">조건 구조화<b>16개</b>완전 자격 검증 아님</div><div class="card">미검토 잠금<b>445개</b>자동 통과하지 않음</div><div class="card">기존 추천 제외<b>453개</b>재포함하지 않음</div></div>
<h2>수정된 처리 방식</h2><ul><li>정책별로 AND/OR/NOT, 조건 주체, 원문, 검토상태를 저장합니다. 검색 태그는 자격 제한으로 쓰지 않습니다.</li>
<li>소득은 구간 경계와 개념을 보존합니다.4인 이상을4인으로 확정하거나 정부지원 생활비를 법정 수급으로 추정하지 않습니다.</li>
<li>관측되지 않은 등록상태·소득인정액·상담/교육 필요성은 미확인으로 남깁니다.</li>
<li>자녀·배우자 정보를 본인의 값으로 대신하지 않습니다. 예외 경로를 전체 연령 범위에 합치지 않습니다.</li>
<li>0/1/2 값을 사람과 정책에서 바로 비교하지 않습니다. 결과는 관측조건 부합·비해당·미확인으로 계산합니다.</li></ul>
<h2>기존8개 분석에 미친 영향</h2><p>새 조건식으로 계산한 기본대상/일반안내 부합 가중평균은 <strong>{current['weighted_mean_positive']:.2f}개</strong>, 미확인 가중평균은 <strong>{current['weighted_mean_unknown']:.2f}개</strong>입니다. 기존6.02개는 상담·디지털 어려움 등의 미관측 조건을 생략한 값이므로 대체했습니다. 두 값 모두 실제 수급 가능한 지원금 개수가 아닙니다.</p>
<p>전체33,944명·18,467가구를 유지했습니다. 모델B는 별도 요건 없는 일반 안내2개만 출력하는 한정된 범위를 유지합니다. 추가 질문 없이 전체 자격을 판정할 수 있다는 뜻은 아닙니다.</p>
<h2>감사 항목별 수정</h2><div class="table">{fixes.to_html(index=False,escape=True,border=0)}</div>
<h2>검증 방법</h2><p>오류가 재현됐던 연령·학생 제외·IQ 누락·퇴직 경로를 합성 사례로 검사했습니다.3값논리의 모든 두 입력 조합, 본인/자녀 주체 분리, 가구 응답 충돌, 결측·구간 보존 및 기존 제외정책 재포함 방지도 검사했습니다. 이는 소프트웨어 검사이며 실제 행정 자격 정확도를 측정한 것이 아닙니다.</p>
<p><a href="tables/common_v3_fixes.csv">수정내역 CSV</a> · <a href="tables/common_v3_results.xlsx">수정 결과·정책 진단 Excel</a> · <a href="reproduce/common_v3_README.md">재현 방법</a></p>
<p>개인 원자료와 개인별 결과는 공개하지 않습니다. 수정 전 코드·보고서는 로컬 Archive에 보존했습니다.</p>'''
    template=(site/'index.html').read_text(encoding='utf-8')
    page=re.sub(r'<main>[\s\S]*?</main>',lambda m:'<main>'+body+'</main>',template,count=1)
    page=re.sub(r'<title>.*?</title>','<title>공통변수 v3 수정 결과</title>',page,count=1)
    (site/'preprocessing_fixes.html').write_text(page,encoding='utf-8')
    replacements={
      '현재 판정에는 국적·복지카드 소유·교육정도·재학상태만 사용합니다.':'v3는 직접 관측된 국적·학교 재학을 사용하고 등록자격·상담·교육 필요성은 미확인으로 남깁니다.',
      '카드 미소유를 미등록장애로 단정하지 않습니다. 장애인 정보화교육에서는 카드 소유 응답에 기본 안내를 부여하고 나머지는 추가확인으로 남겼습니다.':'카드 소유 응답은 원관측으로 보존합니다. 장애인 정보화교육의 등록자격과 교육 필요성은 관측되지 않아 현재 판정은 미확인입니다.',
      '소유1은 장애인 교육 안내에 활용. 미소유2와 결측은 등록상태 미확인':'소유1·미소유2는 카드 응답일 뿐입니다. 법정 등록상태와 교육 필요성은 별도 미관측으로 보존합니다.',
      '두 조건 모두 참이면1; 관측된 필수조건 하나라도 거짓이면0; 나머지는 미확인':'초중고 재학에 부합해도 상담 필요성은 미확인. 관측된 재학 필수조건이 거짓이면0입니다.',
      '1 대한민국이면 국민 대상 일반 안내. 외국국적/결측은 일괄 탈락시키지 않음':'국적은 국민 대상조건의 일부만 확인합니다. 상담·디지털 어려움 등 추가조건은 미확인. 외국국적/결측은 자동 탈락시키지 않습니다.',
      '이용 희망을 전제로 하며, 실제 신청 승인·선정·급여 지급을 예측하지 않습니다.':'미관측 상담·교육 필요성은 미확인으로 남기며, 실제 신청 승인·선정·급여 지급을 예측하지 않습니다.',
      '일반 안내 서비스가 다수여서 평균값이 높게 나타납니다.':'별도 요건 없는 일반 안내와 관측 가능한 조건만 부합으로 집계했습니다.',
      '여기서 기본대상 부합은 저장된 정책 원문의 일반 이용대상 조건과 일치한다는 뜻입니다.':'여기서 기본대상 부합은 v3의 관측조건 또는 명시적 일반안내 범위에 부합한다는 뜻이며 실제 자격 확정이 아닙니다.',
    }
    for p in site.glob('*.html'):
        text=p.read_text(encoding='utf-8')
        for before,after in replacements.items():text=text.replace(before,after)
        if 'href="preprocessing_fixes.html">전처리 수정</a>' not in text:text=text.replace('</nav>','<a href="preprocessing_fixes.html">전처리 수정</a></nav>',1)
        if p.name not in ['preprocessing_fixes.html','eda.html'] and '<!-- common-v3-note -->' not in text:
            text=text.replace('<main>','<main><!-- common-v3-note --><p class="note">공통변수 오류를 수정한 v3 결과입니다. 미관측 조건을 충족으로 처리하지 않습니다. <a href="preprocessing_fixes.html">수정 내역과 검증 범위</a></p>',1)
        p.write_text(text,encoding='utf-8')
    shutil.copy2(OUT/'수정내역18개.csv',site/'tables/common_v3_fixes.csv')
    shutil.copy2(OUT/'공통변수_수정결과.xlsx',site/'tables/common_v3_results.xlsx')
    for name in ['조건판정_v3.py','공통변수_v3.py','test_공통변수_v3.py']:
        shutil.copy2(ROOT/'Code'/name,site/'reproduce'/name)
    (site/'reproduce/common_v3_README.md').write_text('''# 공통변수 v3 재현

기존 정책/사회조사 원자료 및 활성정책 설정을 가진 로컬 프로젝트에서 실행합니다.
이 폴더의 조건판정_v3.py, 공통변수_v3.py, test_공통변수_v3.py를 Code 폴더에 같은 이름으로 복사합니다.
policy_analysis.py는 Code/정책제외_재분석.py, no_question_analysis.py는 Code/무질문모델.py로 배치합니다.
`python Code/정책제외_재분석.py`가 전처리,8개 분석,모델B,수정 보고서를 순서대로 생성합니다.
전처리만 하려면 `python Code/공통변수_v3.py`를 실행합니다.
원자료 배치는 기존 README.md와 no_question_README.md를 참고하세요. pandas,numpy,openpyxl,matplotlib이 필요합니다.
`python -m unittest discover -s Code -p test_공통변수_v3.py -v`로 검사합니다.
구조화16개는 전체자격 검증 완료를 뜻하지 않습니다.445개는 미검토 잠금이며 기존453개 제외는 유지합니다.
정책 원문 해시가 바뀌면 실행을 중단합니다. 개인 원자료는 공개 저장소에 없습니다.
''',encoding='utf-8')
    for name in ['README.md','reproduce/README.md']:
        p=site/name;txt=p.read_text(encoding='utf-8')
        if 'common_v3_README.md' not in txt:p.write_text(txt+'\n공통변수 v3 수정이 반영되었습니다. 추가 실행 파일과 재현 방법은 reproduce/common_v3_README.md를 참고하세요.\n',encoding='utf-8')


if __name__=='__main__':print(json.dumps(build(),ensure_ascii=False,indent=2))
